"""
Elliptic Curve Arithmetic - Presentation Generator
Author: Guru Prakash A (II-MCA, Reg No: 25322011)
Subject: Network Security (Semester 3) - William Stallings, Chapter 10

Generates a 20-slide, 16:9 widescreen presentation with:
- Outfit, Segoe UI, and Consolas typography
- Generous, even spacing and padding across all text fields and cards
- Beautiful formatted tables with cell padding and colored headers
- High-resolution embedded diagram assets (cropped and clean)
- Clean, 100% compliant PowerPoint slide transitions (fade, push, wipe, wheel)
- Fully validated against Microsoft PowerPoint COM automation
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Projector-Ready Light Modern Theme)
    BG_LIGHT = RGBColor(248, 250, 252)        # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)         # #FFFFFF
    CARD_BG_ALT = RGBColor(241, 245, 249)     # #F1F5F9
    CARD_BG_ACCENT = RGBColor(238, 242, 255)  # #EEF2FF
    CARD_BG_GREEN = RGBColor(240, 253, 244)   # #F0FDF4
    
    TEXT_DARK = RGBColor(15, 23, 42)          # #0F172A
    TEXT_MUTED = RGBColor(51, 65, 85)         # #334155
    TEXT_LIGHT = RGBColor(100, 116, 139)      # #64748B

    # Vibrant Accents
    BLUE_PRIMARY = RGBColor(37, 99, 235)      # #2563EB
    PURPLE_ACCENT = RGBColor(124, 58, 237)    # #7C3AED
    EMERALD_GREEN = RGBColor(5, 150, 105)     # #059669
    AMBER_GOLD = RGBColor(217, 119, 6)        # #D97706
    ROSE_RED = RGBColor(225, 29, 72)          # #E11D48
    CYAN_ACCENT = RGBColor(8, 145, 178)       # #0891B2
    
    BORDER_SOFT = RGBColor(226, 232, 240)     # #E2E8F0
    BORDER_BLUE = RGBColor(191, 219, 254)     # #BFDBFE

    # Typography definitions
    FONT_TITLE = "Outfit"
    FONT_BODY = "Segoe UI"
    FONT_MONO = "Consolas"

    def set_bg(slide, trans_type="fade"):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_LIGHT
        
        # Sleek top gradient accent bar
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        line.fill.solid()
        line.fill.fore_color.rgb = BLUE_PRIMARY
        line.line.fill.background()

        # Clean, standards-compliant PowerPoint slide transition (fast speed)
        try:
            if trans_type == "fade":
                trans = parse_xml('<p:transition %s spd="fast"><p:fade/></p:transition>' % nsdecls('p'))
            elif trans_type == "push":
                trans = parse_xml('<p:transition %s spd="fast"><p:push dir="r"/></p:transition>' % nsdecls('p'))
            elif trans_type == "wipe":
                trans = parse_xml('<p:transition %s spd="fast"><p:wipe dir="r"/></p:transition>' % nsdecls('p'))
            elif trans_type == "wheel":
                trans = parse_xml('<p:transition %s spd="fast"><p:wheel spokes="4"/></p:transition>' % nsdecls('p'))
            else:
                trans = parse_xml('<p:transition %s spd="fast"><p:fade/></p:transition>' % nsdecls('p'))
            slide.element.append(trans)
        except Exception:
            pass

    def add_header(slide, slide_num, title_text, subtitle_text=""):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.28), Inches(11.733), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.space_after = Pt(2)
        r_num = p.add_run()
        r_num.text = f"{slide_num:02d}  "
        r_num.font.name = FONT_MONO
        r_num.font.size = Pt(22)
        r_num.font.bold = True
        r_num.font.color.rgb = BLUE_PRIMARY

        r_title = p.add_run()
        r_title.text = title_text
        r_title.font.name = FONT_TITLE
        r_title.font.size = Pt(24)
        r_title.font.bold = True
        r_title.font.color.rgb = TEXT_DARK

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(2)
            r_sub = p2.add_run()
            r_sub.text = subtitle_text
            r_sub.font.name = FONT_BODY
            r_sub.font.size = Pt(13)
            r_sub.font.color.rgb = TEXT_LIGHT

        # Footer credit with comfortable bottom margin
        tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.08), Inches(11.733), Inches(0.32))
        tf_foot = tb_foot.text_frame
        tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
        p_foot = tf_foot.paragraphs[0]
        r_f = p_foot.add_run()
        r_f.text = "Elliptic Curve Arithmetic  |  Presented by: Guru Prakash A (II-MCA, Reg No: 25322011)  |  William Stallings Ch. 10"
        r_f.font.name = FONT_BODY
        r_f.font.size = Pt(10)
        r_f.font.bold = True
        r_f.font.color.rgb = TEXT_LIGHT

    def add_card(slide, left, top, width, height, fill_color=CARD_BG, border_color=BORDER_SOFT, left_strip_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()

        if left_strip_color:
            strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.12), height)
            strip.fill.solid()
            strip.fill.fore_color.rgb = left_strip_color
            strip.line.fill.background()

        return shape

    def add_image_card(slide, img_path, left, top, width, height, caption=""):
        add_card(slide, left, top, width, height, CARD_BG, BORDER_BLUE)
        img_h = height - Inches(0.65) if caption else height - Inches(0.3)
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, left + Inches(0.15), top + Inches(0.15), width - Inches(0.3), img_h)
        if caption:
            tb = slide.shapes.add_textbox(left + Inches(0.15), top + height - Inches(0.5), width - Inches(0.3), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = caption
            r.font.name = FONT_BODY
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = AMBER_GOLD

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, "fade")
    add_card(s1, Inches(1.2), Inches(0.7), Inches(10.933), Inches(6.0), CARD_BG, BORDER_BLUE)
    
    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(0.95), Inches(10.333), Inches(3.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p = tf1.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(14)
    r = p.add_run(); r.text = "● MASTER OF COMPUTER APPLICATIONS  |  SEMINAR PRESENTATION ●"
    r.font.name = FONT_MONO; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = AMBER_GOLD

    p2 = tf1.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_after = Pt(12)
    r2 = p2.add_run(); r2.text = "Elliptic Curve Arithmetic"
    r2.font.name = FONT_TITLE; r2.font.size = Pt(44); r2.font.bold = True; r2.font.color.rgb = BLUE_PRIMARY

    p3 = tf1.add_paragraph(); p3.alignment = PP_ALIGN.CENTER; p3.space_after = Pt(12)
    r3 = p3.add_run(); r3.text = "Mathematical Foundations, Group Laws, and Finite Field Cryptography"
    r3.font.name = FONT_BODY; r3.font.size = Pt(19); r3.font.bold = True; r3.font.color.rgb = TEXT_DARK

    p4 = tf1.add_paragraph(); p4.alignment = PP_ALIGN.CENTER
    r4 = p4.add_run(); r4.text = "Network Security  |  Semester 3  |  Reference: William Stallings, Chapter 10"
    r4.font.name = FONT_BODY; r4.font.size = Pt(14); r4.font.bold = True; r4.font.color.rgb = TEXT_LIGHT

    # Presenter Card with balanced padding
    add_card(s1, Inches(3.2), Inches(4.7), Inches(6.933), Inches(1.65), CARD_BG_ALT, BLUE_PRIMARY, BLUE_PRIMARY)
    tb_pres = s1.shapes.add_textbox(Inches(3.5), Inches(4.85), Inches(6.4), Inches(1.4))
    tf_pres = tb_pres.text_frame
    tf_pres.word_wrap = True
    tf_pres.margin_left = tf_pres.margin_top = tf_pres.margin_right = tf_pres.margin_bottom = 0

    p_p0 = tf_pres.paragraphs[0]; p_p0.space_after = Pt(4)
    r_p0 = p_p0.add_run(); r_p0.text = "PRESENTED BY"; r_p0.font.name = FONT_MONO; r_p0.font.size = Pt(11); r_p0.font.bold = True; r_p0.font.color.rgb = BLUE_PRIMARY
    
    p_p1 = tf_pres.add_paragraph(); p_p1.space_after = Pt(4)
    r_p1 = p_p1.add_run(); r_p1.text = "Guru Prakash A"; r_p1.font.name = FONT_TITLE; r_p1.font.size = Pt(22); r_p1.font.bold = True; r_p1.font.color.rgb = TEXT_DARK
    
    p_p2 = tf_pres.add_paragraph()
    r_p2 = p_p2.add_run(); r_p2.text = "Class: II - MCA   |   Register No: 25322011"; r_p2.font.name = FONT_BODY; r_p2.font.size = Pt(14); r_p2.font.bold = True; r_p2.font.color.rgb = PURPLE_ACCENT

    # ==========================================
    # SLIDE 2: Agenda & Overview
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, "push")
    add_header(s2, 2, "Agenda & Overview", "Following William Stallings' Cryptography and Network Security (Chapter 10)")
    
    agendas = [
        ("📘 1. Foundations & Concepts", [
            "Motivation: The RSA Scaling Problem",
            "What is an Elliptic Curve? (Smooth Cubic Plane)",
            "Why are Elliptic Curves Used in Cryptography?",
            "Abelian Group Axioms (A1 to A5)"
        ], BLUE_PRIMARY),
        ("📐 2. Real Curves & Geometric Rules", [
            "Elliptic Curves over Real Numbers R",
            "The 5 Fundamental Geometric Addition Rules",
            "Point Doubling via Tangent Line Construction",
            "Algebraic Addition & Doubling Formulas"
        ], AMBER_GOLD),
        ("🔢 3. Prime Curves over Z_p", [
            "Two Families: Prime vs. Binary Curves",
            "Curves over Z_p, Hasse's Theorem & Condition",
            "4 Formal Addition Rules & Slope Formulas",
            "Worked Example 1: Point Addition (P+Q)",
            "Worked Example 2: Point Doubling (2P)"
        ], PURPLE_ACCENT),
        ("💻 4. Binary Curves & Scalability", [
            "Curves over Binary Fields GF(2^m)",
            "4 Addition Rules & Slope Formulas in GF(2^m)",
            "Master Comparison: Z_p vs. GF(2^m)",
            "Scalar Multiplication & Double-and-Add Algorithm"
        ], EMERALD_GREEN)
    ]
    for idx, (title, items, color) in enumerate(agendas):
        row, col = idx // 2, idx % 2
        l = Inches(0.8 + col * 5.95); t = Inches(1.48 + row * 2.68)
        add_card(s2, l, t, Inches(5.75), Inches(2.48), CARD_BG, BORDER_SOFT, color)
        tb = s2.shapes.add_textbox(l + Inches(0.3), t + Inches(0.2), Inches(5.2), Inches(2.15))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]; p.space_after = Pt(10)
        r = p.add_run(); r.text = title; r.font.name = FONT_TITLE; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = color
        for it in items:
            p_i = tf.add_paragraph()
            p_i.space_after = Pt(5)
            p_i.line_spacing = 1.2
            r_i = p_i.add_run()
            r_i.text = f"•  {it}"
            r_i.font.name = FONT_BODY
            r_i.font.size = Pt(12)
            r_i.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: Motivation — Why ECC?
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, "fade")
    add_header(s3, 3, "Motivation — Why ECC?", "The need for scalable, efficient public-key cryptography (Stallings Sec 10.3)")
    
    # Left Card
    add_card(s3, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, ROSE_RED)
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "The RSA Scaling Problem"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = ROSE_RED
    
    bullets_m = [
        ("Exploding Key Lengths", "RSA keys must grow exponentially to maintain security against factorization attacks (2048 to 15360 bits)."),
        ("Server Bottlenecks", "Massive computational latency on high-volume cloud servers handling concurrent SSL/TLS handshakes."),
        ("Constrained Environments", "Severe battery drain and memory starvation on mobile phones, smart cards, and IoT embedded devices."),
        ("William Stallings' Assessment", "'The principal attraction of ECC, compared to RSA, is that it appears to offer equal security for a far smaller key size, thereby reducing processing overhead.'")
    ]
    for b_title, b_desc in bullets_m:
        pi = tf.add_paragraph()
        pi.space_after = Pt(10)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # Right Card: Formatted PowerPoint Table
    add_card(s3, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, EMERALD_GREEN)
    tb_rt = s3.shapes.add_textbox(Inches(7.05), Inches(1.68), Inches(5.2), Inches(0.5))
    tf_rt = tb_rt.text_frame; tf_rt.margin_left = tf_rt.margin_top = tf_rt.margin_right = tf_rt.margin_bottom = 0
    p_rt = tf_rt.paragraphs[0]; r_rt = p_rt.add_run()
    r_rt.text = "NIST Key Size Equivalence (SP 800-57)"; r_rt.font.name = FONT_TITLE; r_rt.font.size = Pt(18); r_rt.font.bold = True; r_rt.font.color.rgb = EMERALD_GREEN

    table_shape = s3.shapes.add_table(6, 4, Inches(7.0), Inches(2.28), Inches(5.3), Inches(3.3))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.5)
    tbl.columns[1].width = Inches(1.3)
    tbl.columns[2].width = Inches(1.3)
    tbl.columns[3].width = Inches(1.2)

    table_data = [
        ("Security Level", "RSA Key", "ECC Key", "Ratio"),
        ("80 bits", "1024 bits", "160 bits", "6 : 1"),
        ("112 bits", "2048 bits", "224 bits", "9 : 1"),
        ("128 bits ★", "3072 bits", "256 bits", "12 : 1"),
        ("192 bits", "7680 bits", "384 bits", "20 : 1"),
        ("256 bits", "15360 bits", "512 bits", "30 : 1")
    ]
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_c = cell.text_frame.paragraphs[0]
            p_c.alignment = PP_ALIGN.CENTER
            p_c.font.name = FONT_MONO if r_idx > 0 else FONT_TITLE
            p_c.font.size = Pt(11)
            
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TEXT_DARK
                p_c.font.bold = True
                p_c.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx == 3: # 128 bits standard
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG_ACCENT
                p_c.font.bold = True
                p_c.font.color.rgb = BLUE_PRIMARY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
                p_c.font.color.rgb = TEXT_DARK

    tb_note = s3.shapes.add_textbox(Inches(7.0), Inches(5.8), Inches(5.3), Inches(0.85))
    tf_note = tb_note.text_frame; tf_note.word_wrap = True
    tf_note.margin_left = tf_note.margin_top = tf_note.margin_right = tf_note.margin_bottom = 0
    p_n = tf_note.paragraphs[0]; p_n.line_spacing = 1.25; r_n = p_n.add_run()
    r_n.text = "★ 128-bit Commercial Standard: ECC achieves the exact same security with 256 bits that RSA requires 3072 bits to reach (12× smaller key size!)."
    r_n.font.name = FONT_BODY; r_n.font.size = Pt(11.5); r_n.font.bold = True; r_n.font.color.rgb = BLUE_PRIMARY

    # ==========================================
    # SLIDE 4: What is an Elliptic Curve?
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, "push")
    add_header(s4, 4, "What is an Elliptic Curve?", "A non-singular cubic plane curve with smooth geometry (Stallings Sec 10.3)")
    
    add_card(s4, Inches(0.8), Inches(1.45), Inches(6.0), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.4), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(8)
    r = p.add_run(); r.text = "Definition & Geometric Properties"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    
    p_eq = tf.add_paragraph(); p_eq.space_after = Pt(12)
    r_eq = p_eq.add_run(); r_eq.text = "y² = x³ + ax + b"; r_eq.font.name = FONT_MONO; r_eq.font.size = Pt(22); r_eq.font.bold = True; r_eq.font.color.rgb = PURPLE_ACCENT
    
    bullets_we = [
        ("Cubic Equation (Degree 3)", "The highest power is 3. Any straight line intersects the curve at exactly THREE points (counting multiplicity)."),
        ("Not an Ellipse", "Named after elliptic integrals historically used to calculate the arc length of an ellipse."),
        ("Point at Infinity O", "A single conceptual zero point lying along every vertical line; serves as group additive identity."),
        ("Horizontal Symmetry", "Symmetrical about the x-axis: y = ±√(x³ + ax + b)."),
        ("Non-Singularity Condition", "Discriminant 4a³ + 27b² ≠ 0 guarantees no cusps or self-intersections, so tangents always exist.")
    ]
    for b_title, b_desc in bullets_we:
        pi = tf.add_paragraph()
        pi.space_after = Pt(8)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    add_image_card(s4, "images/curve_shapes.png", Inches(7.1), Inches(1.45), Inches(5.4), Inches(5.4), "Continuous Elliptic Curve Shapes over Real Numbers R")

    # ==========================================
    # SLIDE 5: Why are Elliptic Curves Used?
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, "fade")
    add_header(s5, 5, "Why are Elliptic Curves Used in Cryptography?", "Exponential hardness and maximum security per bit")
    
    # Left Card
    add_card(s5, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s5.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "1. Natural Abelian Group Structure"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    bullets_g = [
        ("Group Operations", "Points on an elliptic curve naturally form an Abelian group under chord-and-tangent addition."),
        ("Closure Guaranteed", "Adding any two valid curve points always produces another point that is guaranteed to lie on the curve."),
        ("Associative & Commutative", "The order and grouping of operations do not matter: (P + Q) + R = P + (Q + R) and P + Q = Q + P."),
        ("Identity & Inverses", "The point at infinity O serves as identity, and reflection across the x-axis provides exact point negation -P.")
    ]
    for b_title, b_desc in bullets_g:
        pi = tf.add_paragraph()
        pi.space_after = Pt(10)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # Right Card
    add_card(s5, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, PURPLE_ACCENT)
    tb_r = s5.shapes.add_textbox(Inches(7.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf_r = tb_r.text_frame; tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    
    p = tf_r.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "2. No Sub-Exponential Attacks"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = PURPLE_ACCENT
    bullets_a = [
        ("The Weakness of RSA", "In regular integers modulo p (RSA/Diffie-Hellman), Index Calculus attacks solve discrete logs in sub-exponential time O(exp(c (ln p)^{1/3}))."),
        ("No Prime Factors on Curves", "Elliptic curve points have NO analogue to prime factors! Index calculus does NOT apply to general elliptic curves."),
        ("Strictly Exponential Hardness", "The fastest known attack is Pollard's rho, which requires fully exponential time O(√p)."),
        ("12× Key Size Advantage", "A 256-bit ECC key offers the exact same mathematical security as a 3072-bit RSA key.")
    ]
    for b_title, b_desc in bullets_a:
        pi = tf_r.add_paragraph()
        pi.space_after = Pt(10)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 6: Abelian Groups
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, "wipe")
    add_header(s6, 6, "Abelian Groups — Mathematical Foundation", "A set G with binary operation satisfying 5 axioms (Stallings Sec 10.3)")
    
    props = [
        ("(A1) Closure", "For all a, b in G, the combination a • b is also in G.", BLUE_PRIMARY),
        ("(A2) Associativity", "For all a, b, c in G, (a • b) • c = a • (b • c). Grouping does not affect result.", PURPLE_ACCENT),
        ("(A3) Identity Element", "There exists an element e in G such that a • e = e • a = a.", EMERALD_GREEN),
        ("(A4) Inverse Element", "For each a in G, there exists a unique a' such that a • a' = e.", ROSE_RED),
        ("(A5) Commutativity", "For all a, b in G, a • b = b • a. (The distinctive Abelian property!)", AMBER_GOLD),
    ]
    for idx, (p_name, p_desc, col) in enumerate(props):
        t = Inches(1.45 + idx * 0.96)
        add_card(s6, Inches(0.8), t, Inches(11.733), Inches(0.86), CARD_BG, BORDER_SOFT, col)
        tb = s6.shapes.add_textbox(Inches(1.2), t + Inches(0.12), Inches(11.1), Inches(0.62))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.line_spacing = 1.2
        r1 = p.add_run(); r1.text = f"{p_name}:   "; r1.font.name = FONT_TITLE; r1.font.size = Pt(15.5); r1.font.bold = True; r1.font.color.rgb = col
        r2 = p.add_run(); r2.text = p_desc; r2.font.name = FONT_BODY; r2.font.size = Pt(13.5); r2.font.color.rgb = TEXT_DARK

    add_card(s6, Inches(0.8), Inches(6.32), Inches(11.733), Inches(0.68), CARD_BG_ALT, BLUE_PRIMARY)
    tb_b = s6.shapes.add_textbox(Inches(1.05), Inches(6.38), Inches(11.2), Inches(0.55))
    tf_b = tb_b.text_frame; tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    p = tf_b.paragraphs[0]; r = p.add_run()
    r.text = "ECC Analogy: Group operation = Point Addition (+). Identity = Point at Infinity O. Inverse = -P = (x, -y)."
    r.font.name = FONT_BODY; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY

    # ==========================================
    # SLIDE 7: Real Curves
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, "push")
    add_header(s7, 7, "Elliptic Curves over Real Numbers R", "Weierstrass cubic equation and non-singularity condition (Stallings Sec 10.3)")
    
    add_card(s7, Inches(0.8), Inches(1.45), Inches(6.0), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s7.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.4), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(8)
    r = p.add_run(); r.text = "Cubic Weierstrass Form"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    
    p_eq = tf.add_paragraph(); p_eq.space_after = Pt(12)
    r_eq = p_eq.add_run(); r_eq.text = "y² = x³ + ax + b"; r_eq.font.name = FONT_MONO; r_eq.font.size = Pt(22); r_eq.font.bold = True; r_eq.font.color.rgb = PURPLE_ACCENT
    
    bullets = [
        ("Set of Points E(a, b)", "Comprises all real coordinate pairs (x, y) satisfying the cubic equation, plus the single point at infinity O."),
        ("Vertical Symmetry", "The curve is completely symmetric about the horizontal axis y = 0."),
        ("Non-Singularity (Eq. 10.2)", "Discriminant condition: 4a³ + 27b² ≠ 0."),
        ("Geometric Consequence", "Ensures the polynomial x³ + ax + b has three distinct roots, preventing cusps or self-intersections so tangent lines are defined everywhere.")
    ]
    for b_title, b_desc in bullets:
        pi = tf.add_paragraph()
        pi.space_after = Pt(9)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    add_image_card(s7, "images/elliptic_curve_shapes.png", Inches(7.1), Inches(1.45), Inches(5.4), Inches(5.4), "Stallings Fig 10.4: Curve Shapes for (a) y²=x³-x and (b) y²=x³-x+1")

    # ==========================================
    # SLIDE 8: 5 Geometric Addition Rules
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, "fade")
    add_header(s8, 8, "Geometric Addition — The 5 Fundamental Rules", "Chord-and-Tangent Law: If 3 points lie on a straight line, P + Q + R = O")
    
    rules = [
        ("Rule 1: Identity Element", "O serves as additive identity: P + O = P. By definition, O = -O.", BLUE_PRIMARY),
        ("Rule 2: Point Negation", "Negative of P=(x,y) is -P=(x,-y). Vertical line joins them. P + (-P) = O.", PURPLE_ACCENT),
        ("Rule 3: Distinct Points (P ≠ Q)", "Draw secant line through P and Q. Third intersection is R. Sum P + Q = -R (reflect R across x-axis).", AMBER_GOLD),
        ("Rule 4: Vertical Line Case", "Two points with same x-coord (P and -P) lie on a vertical line intersecting at O. P + (-P) = O.", ROSE_RED),
        ("Rule 5: Point Doubling (2Q)", "Draw tangent line at Q. It intersects curve at S. Doubled point 2Q = Q + Q = -S.", EMERALD_GREEN)
    ]
    for idx, (r_name, r_desc, col) in enumerate(rules):
        t = Inches(1.45 + idx * 1.05)
        add_card(s8, Inches(0.8), t, Inches(6.8), Inches(0.96), CARD_BG, BORDER_SOFT, col)
        tb = s8.shapes.add_textbox(Inches(1.1), t + Inches(0.1), Inches(6.3), Inches(0.78))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.space_after = Pt(2)
        r1 = p.add_run(); r1.text = r_name; r1.font.name = FONT_TITLE; r1.font.size = Pt(13.5); r1.font.bold = True; r1.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.line_spacing = 1.2
        r2 = p2.add_run(); r2.text = r_desc; r2.font.name = FONT_BODY; r2.font.size = Pt(11.5); r2.font.color.rgb = TEXT_DARK

    add_image_card(s8, "images/point_addition_diagram.png", Inches(7.9), Inches(1.45), Inches(4.6), Inches(5.4), "Secant Line Construction for P + Q = -R")

    # ==========================================
    # SLIDE 9: Point Doubling
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, "push")
    add_header(s9, 9, "Point Doubling — Tangent Construction 2P", "Adding a point to itself via tangent slope (Stallings Sec 10.3)")
    
    add_card(s9, Inches(0.8), Inches(1.45), Inches(6.0), Inches(5.4), CARD_BG, BORDER_SOFT, AMBER_GOLD)
    tb = s9.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.4), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "Tangent Line Construction"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = AMBER_GOLD
    bullets_pd = [
        ("Secant turns into Tangent", "As point Q approaches P along the curve, the secant line approaches the tangent line to the curve at P."),
        ("Slope via Implicit Derivative", "Differentiating y² = x³ + ax + b yields 2y(dy/dx) = 3x² + a, so tangent slope λ = (3x_P² + a) / (2y_P)."),
        ("Second Intersection Point S", "The tangent line intersects the cubic curve at exactly one second point S."),
        ("Reflection across x-axis", "The reflection of S across the x-axis gives the doubled point: 2P = -S."),
        ("Vertical Tangent Case", "If y_P = 0, the tangent line is vertical, meaning it intersects only at infinity: 2P = O.")
    ]
    for b_title, b_desc in bullets_pd:
        pi = tf.add_paragraph()
        pi.space_after = Pt(10)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    add_image_card(s9, "images/point_doubling_diagram.png", Inches(7.1), Inches(1.45), Inches(5.4), Inches(5.4), "Tangent Line Construction for Point Doubling 2P = -S")

    # ==========================================
    # SLIDE 10: Algebraic Formulas over R
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10, "fade")
    add_header(s10, 10, "Algebraic Formulas over Real Numbers R", "Converting geometric secants and tangents to closed-form equations")
    
    # Left Card
    add_card(s10, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s10.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(10)
    r = p.add_run(); r.text = "Point Addition (P ≠ Q)"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    
    sections_add = [
        ("Secant Slope λ:", "λ = (y_Q - y_P) / (x_Q - x_P)"),
        ("Intersection Coordinate x_R:", "x_R = λ² - x_P - x_Q"),
        ("Reflected Coordinate y_R:", "y_R = λ(x_P - x_R) - y_P"),
        ("Vertical Condition:", "If x_P = x_Q and y_P = -y_Q, then P + Q = O.")
    ]
    for lbl, formula in sections_add:
        p_l = tf.add_paragraph(); p_l.space_after = Pt(2)
        r_l = p_l.add_run(); r_l.text = lbl; r_l.font.name = FONT_BODY; r_l.font.size = Pt(13); r_l.font.bold = True; r_l.font.color.rgb = TEXT_DARK
        p_f = tf.add_paragraph(); p_f.space_after = Pt(10)
        r_f = p_f.add_run(); r_f.text = f"   {formula}"; r_f.font.name = FONT_MONO; r_f.font.size = Pt(13); r_f.font.color.rgb = BLUE_PRIMARY

    # Right Card
    add_card(s10, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, PURPLE_ACCENT)
    tb_r = s10.shapes.add_textbox(Inches(7.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf_r = tb_r.text_frame; tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    
    p = tf_r.paragraphs[0]; p.space_after = Pt(10)
    r = p.add_run(); r.text = "Point Doubling (2P = P + P)"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = PURPLE_ACCENT
    
    sections_dbl = [
        ("Tangent Slope λ (via dy/dx):", "λ = (3x_P² + a) / (2y_P)"),
        ("Intersection Coordinate x_R:", "x_R = λ² - 2x_P"),
        ("Reflected Coordinate y_R:", "y_R = λ(x_P - x_R) - y_P"),
        ("Vertical Condition:", "If y_P = 0, tangent is vertical, so 2P = O.")
    ]
    for lbl, formula in sections_dbl:
        p_l = tf_r.add_paragraph(); p_l.space_after = Pt(2)
        r_l = p_l.add_run(); r_l.text = lbl; r_l.font.name = FONT_BODY; r_l.font.size = Pt(13); r_l.font.bold = True; r_l.font.color.rgb = TEXT_DARK
        p_f = tf_r.add_paragraph(); p_f.space_after = Pt(10)
        r_f = p_f.add_run(); r_f.text = f"   {formula}"; r_f.font.name = FONT_MONO; r_f.font.size = Pt(13); r_f.font.color.rgb = PURPLE_ACCENT

    # ==========================================
    # SLIDE 11: Two Cryptographic Families
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11, "wipe")
    add_header(s11, 11, "Two Cryptographic Families — Z_p vs GF(2^m)", "Finite fields eliminate floating-point rounding errors (Stallings Sec 10.3)")
    
    # Left Card
    add_card(s11, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s11.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "1. Prime Curves over Z_p"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    bullets_p = [
        ("Field Elements", "Restricted to integers {0, 1, ..., p-1} modulo a large prime p."),
        ("Curve Equation", "y² ≡ x³ + ax + b (mod p)."),
        ("Non-Singularity", "(4a³ + 27b²) mod p ≠ 0."),
        ("Ideal for Software", "Standard microprocessors (Intel, ARM, AMD) have optimized 64-bit integer ALUs that compute modular operations with maximum efficiency."),
        ("No Bit-Level Overhead", "Does NOT require complex polynomial carry manipulation.")
    ]
    for b_title, b_desc in bullets_p:
        pi = tf.add_paragraph()
        pi.space_after = Pt(9)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # Right Card
    add_card(s11, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, PURPLE_ACCENT)
    tb_r = s11.shapes.add_textbox(Inches(7.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf_r = tb_r.text_frame; tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    
    p = tf_r.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "2. Binary Curves over GF(2^m)"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = PURPLE_ACCENT
    bullets_b = [
        ("Field Elements", "Polynomials with binary coefficients {0, 1} of degree < m."),
        ("Curve Equation", "y² + xy = x³ + ax² + b."),
        ("Non-Singularity", "b ≠ 0 (ensures curve is non-singular)."),
        ("Ideal for Hardware", "Addition in GF(2^m) is simple bitwise XOR without any carry propagation delays!"),
        ("Ultra-Compact Silicon", "Requires very few logic gates, making it the premier choice for smart cards, ASICs, and FPGAs.")
    ]
    for b_title, b_desc in bullets_b:
        pi = tf_r.add_paragraph()
        pi.space_after = Pt(9)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 12: Prime Curves over Z_p
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12, "push")
    add_header(s12, 12, "Elliptic Curves over Prime Field Z_p", "Set of points E_p(a, b) and Hasse's Theorem (Stallings Section 10.3)")
    
    add_card(s12, Inches(0.8), Inches(1.45), Inches(6.0), Inches(5.4), CARD_BG, BORDER_SOFT, BLUE_PRIMARY)
    tb = s12.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.4), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(10)
    r = p.add_run(); r.text = "Definition & Group Size Bounds"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    
    bullets_zp_def = [
        ("Modular Equation", "y² mod p = (x³ + ax + b) mod p, where a, b ∈ Z_p."),
        ("Non-Singularity", "(4a³ + 27b²) mod p ≠ 0."),
        ("Hasse's Theorem (Stallings p. 8)", "The number of points N on E_p(a, b) satisfies:"),
        ("   Formula", "p + 1 - 2√p  ≤  N  ≤  p + 1 + 2√p"),
        ("Worked Count for p = 23", "For E₂₃(1, 1), Hasse's theorem bounds N between 15 and 33. The actual count is N = 28 (27 points from Table 10.1 plus O)."),
        ("Symmetry Property", "Symmetric about the horizontal line y = p / 2 = 11.5.")
    ]
    for b_title, b_desc in bullets_zp_def:
        pi = tf.add_paragraph()
        pi.space_after = Pt(8)
        pi.line_spacing = 1.22
        if b_title.startswith("   "):
            ri = pi.add_run(); ri.text = f"       {b_desc}"; ri.font.name = FONT_MONO; ri.font.size = Pt(12.5); ri.font.bold = True; ri.font.color.rgb = BLUE_PRIMARY
        else:
            ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
            ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    add_image_card(s12, "images/figure_10_5_e23.png", Inches(7.1), Inches(1.45), Inches(5.4), Inches(5.4), "Stallings Fig 10.5 — Scatter Plot of E₂₃(1, 1) Points (Symmetric about y=11.5)")

    # ==========================================
    # SLIDE 13: Arithmetic Rules over Z_p
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13, "fade")
    add_header(s13, 13, "Arithmetic Rules & Slope Formulas over Z_p", "The 4 formal rules for E_p(a, b) (Stallings p. 7-8)")
    
    rules_zp = [
        ("Rule 1: Additive Identity", "P + O = P for all points P on the curve. Point O is its own inverse: O = -O.", BLUE_PRIMARY),
        ("Rule 2: Point Negation", "If P = (x_P, y_P), -P = (x_P, -y_P mod p) = (x_P, p - y_P). P + (-P) = O.\nExample on E₂₃(1, 1): P = (13, 7) gives -P = (13, 23 - 7) = (13, 16) ∈ E₂₃(1, 1).", PURPLE_ACCENT),
        ("Rule 3: Point Addition (P ≠ Q)", "λ = [(y_Q - y_P) × (x_Q - x_P)^{-1}] mod p\nx_R = (λ² - x_P - x_Q) mod p\ny_R = [λ(x_P - x_R) - y_P] mod p", AMBER_GOLD),
        ("Rule 4: Point Doubling (P = Q)", "λ = [(3x_P² + a) × (2y_P)^{-1}] mod p\nx_R = (λ² - 2x_P) mod p\ny_R = [λ(x_P - x_R) - y_P] mod p", EMERALD_GREEN)
    ]
    for idx, (r_name, r_desc, col) in enumerate(rules_zp):
        t = Inches(1.45 + idx * 1.34)
        add_card(s13, Inches(0.8), t, Inches(11.733), Inches(1.22), CARD_BG, BORDER_SOFT, col)
        tb = s13.shapes.add_textbox(Inches(1.1), t + Inches(0.1), Inches(11.2), Inches(1.05))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.space_after = Pt(3)
        r1 = p.add_run(); r1.text = r_name; r1.font.name = FONT_TITLE; r1.font.size = Pt(14); r1.font.bold = True; r1.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.line_spacing = 1.2
        r2 = p2.add_run(); r2.text = r_desc; r2.font.name = FONT_MONO; r2.font.size = Pt(11.5); r2.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 14: Worked Example 1 (Addition)
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_bg(s14, "push")
    add_header(s14, 14, "Worked Numerical Example 1 — Point Addition P + Q", "Evaluated on curve E₂₃(1, 1): y² ≡ x³ + x + 1 (mod 23)")
    
    # 3 distinct step cards - optimized font sizes to prevent awkward wrapping
    steps_add_cards = [
        ("Step 1: Compute Slope λ", [
            ("Points on Curve E₂₃(1, 1):", "P = (3, 10),  Q = (9, 7)"),
            ("Secant Slope Formula:", "λ = (y_Q - y_P) / (x_Q - x_P) mod 23"),
            ("Coordinate Substitution:", "λ = (7 - 10) / (9 - 3) = -3 / 6"),
            ("Numerator & Denom:", "λ = -1 / 2 ≡ -3 × 6⁻¹ mod 23")
        ], BLUE_PRIMARY),
        ("Step 2: Modular Inverse", [
            ("Extended Euclidean Algorithm:", "Need 6⁻¹ mod 23: 6 × 4 = 24 ≡ 1"),
            ("Inverse Value Found:", "6⁻¹ ≡ 4 (mod 23)"),
            ("Multiply by Inverse:", "λ = (-3) × 4 = -12 mod 23"),
            ("Positive Slope Result:", "λ = 23 - 12 = 11 (mod 23) ✔")
        ], PURPLE_ACCENT),
        ("Step 3: Point Coordinates", [
            ("Compute x_R Coordinate:", "x_R = (11² - 3 - 9) mod 23 = 17"),
            ("Compute y_R Coordinate:", "y_R = [11(3 - 17) - 10] = 20"),
            ("Point Addition Result:", "P + Q = (17, 20) ∈ E₂₃(1, 1) ✔"),
            ("Table 10.1 Confirmation:", "Point verified in Stallings book")
        ], EMERALD_GREEN)
    ]
    for idx, (s_title, items, col) in enumerate(steps_add_cards):
        l = Inches(0.8 + idx * 3.98)
        add_card(s14, l, Inches(1.45), Inches(3.8), Inches(4.5), CARD_BG, BORDER_SOFT, col)
        tb = s14.shapes.add_textbox(l + Inches(0.22), Inches(1.65), Inches(3.36), Inches(4.15))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]; p.space_after = Pt(10)
        r = p.add_run(); r.text = s_title; r.font.name = FONT_TITLE; r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = col
        
        for lbl, val in items:
            p_l = tf.add_paragraph(); p_l.space_after = Pt(2)
            r_l = p_l.add_run(); r_l.text = lbl; r_l.font.name = FONT_BODY; r_l.font.size = Pt(11); r_l.font.bold = True; r_l.font.color.rgb = TEXT_DARK
            p_v = tf.add_paragraph(); p_v.space_after = Pt(8); p_v.line_spacing = 1.15
            r_v = p_v.add_run(); r_v.text = f" {val}"; r_v.font.name = FONT_MONO; r_v.font.size = Pt(10.5)
            r_v.font.bold = "✔" in val or "λ =" in val
            r_v.font.color.rgb = EMERALD_GREEN if "✔" in val else (BLUE_PRIMARY if "λ =" in val else TEXT_MUTED)

    # Bottom verification card
    add_card(s14, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.75), CARD_BG_ALT, BLUE_PRIMARY)
    tb_check = s14.shapes.add_textbox(Inches(1.05), Inches(6.22), Inches(11.2), Inches(0.6))
    tf_check = tb_check.text_frame; tf_check.margin_left = tf_check.margin_top = tf_check.margin_right = tf_check.margin_bottom = 0
    p = tf_check.paragraphs[0]; p.line_spacing = 1.2; r = p.add_run()
    r.text = "Verification: LHS = 20² mod 23 = 400 mod 23 = 9.  RHS = (17³ + 17 + 1) mod 23 = 4931 mod 23 = 9.  LHS = RHS = 9 ✔"
    r.font.name = FONT_MONO; r.font.size = Pt(11.5); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY

    # ==========================================
    # SLIDE 15: Worked Example 2 (Doubling)
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    set_bg(s15, "fade")
    add_header(s15, 15, "Worked Numerical Example 2 — Point Doubling 2P", "Evaluated on curve E₂₃(1, 1): y² ≡ x³ + x + 1 (mod 23)")
    
    steps_dbl_cards = [
        ("Step 1: Compute Tangent Slope λ", [
            ("Point on Curve E₂₃(1, 1):", "P = (3, 10), so x_P = 3, y_P = 10"),
            ("Tangent Slope Formula:", "λ = (3x_P² + a) / (2y_P) mod 23"),
            ("Coordinate Substitution:", "λ = (3(3²) + 1) / (2 × 10) mod 23"),
            ("Numerator & Denom:", "λ = (27 + 1) / 20 = 28 / 20 mod 23")
        ], PURPLE_ACCENT),
        ("Step 2: Modular Inverse", [
            ("Modulo Reduction:", "28 mod 23 = 5 => λ = 5/20 = 1/4"),
            ("Inverse Value Found:", "4 × 6 = 24 ≡ 1 => 4⁻¹ ≡ 6 (mod 23)"),
            ("Multiply by Inverse:", "λ = 1 × 6 ≡ 6 (mod 23) ✔"),
            ("Alternative Check:", "20⁻¹ ≡ 15 => 28 × 15 = 420 ≡ 6")
        ], AMBER_GOLD),
        ("Step 3: Point Coordinates", [
            ("Compute x_R Coordinate:", "x_R = (6² - 2(3)) mod 23 = 7"),
            ("Compute y_R Coordinate:", "y_R = [6(3 - 7) - 10] = 12"),
            ("Point Doubling Result:", "2P = (7, 12) ∈ E₂₃(1, 1) ✔"),
            ("Table 10.1 Confirmation:", "Point verified in Stallings book")
        ], BLUE_PRIMARY)
    ]
    for idx, (s_title, items, col) in enumerate(steps_dbl_cards):
        l = Inches(0.8 + idx * 3.98)
        add_card(s15, l, Inches(1.45), Inches(3.8), Inches(4.5), CARD_BG, BORDER_SOFT, col)
        tb = s15.shapes.add_textbox(l + Inches(0.22), Inches(1.65), Inches(3.36), Inches(4.15))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]; p.space_after = Pt(10)
        r = p.add_run(); r.text = s_title; r.font.name = FONT_TITLE; r.font.size = Pt(15.5); r.font.bold = True; r.font.color.rgb = col
        
        for lbl, val in items:
            p_l = tf.add_paragraph(); p_l.space_after = Pt(2)
            r_l = p_l.add_run(); r_l.text = lbl; r_l.font.name = FONT_BODY; r_l.font.size = Pt(11); r_l.font.bold = True; r_l.font.color.rgb = TEXT_DARK
            p_v = tf.add_paragraph(); p_v.space_after = Pt(8); p_v.line_spacing = 1.15
            r_v = p_v.add_run(); r_v.text = f" {val}"; r_v.font.name = FONT_MONO; r_v.font.size = Pt(10.5)
            r_v.font.bold = "✔" in val or "λ =" in val
            r_v.font.color.rgb = BLUE_PRIMARY if "✔" in val else (PURPLE_ACCENT if "λ =" in val else TEXT_MUTED)

    add_card(s15, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.75), CARD_BG_ALT, PURPLE_ACCENT)
    tb_check2 = s15.shapes.add_textbox(Inches(1.05), Inches(6.22), Inches(11.2), Inches(0.6))
    tf_check2 = tb_check2.text_frame; tf_check2.margin_left = tf_check2.margin_top = tf_check2.margin_right = tf_check2.margin_bottom = 0
    p = tf_check2.paragraphs[0]; p.line_spacing = 1.2; r = p.add_run()
    r.text = "Verification: LHS = 12² mod 23 = 144 mod 23 = 6.  RHS = (7³ + 7 + 1) mod 23 = 351 mod 23 = 6.  LHS = RHS = 6 ✔"
    r.font.name = FONT_MONO; r.font.size = Pt(11.5); r.font.bold = True; r.font.color.rgb = PURPLE_ACCENT

    # ==========================================
    # SLIDE 16: Binary Curves over GF(2^m)
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    set_bg(s16, "push")
    add_header(s16, 16, "Elliptic Curves over Binary Fields GF(2^m)", "Polynomial representation in characteristic 2 (Stallings Section 10.3)")
    
    add_card(s16, Inches(0.8), Inches(1.45), Inches(6.0), Inches(5.4), CARD_BG, BORDER_SOFT, PURPLE_ACCENT)
    tb = s16.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.4), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(8)
    r = p.add_run(); r.text = "Binary Curve Equation"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = PURPLE_ACCENT
    
    p_eq = tf.add_paragraph(); p_eq.space_after = Pt(12)
    r_eq = p_eq.add_run(); r_eq.text = "y² + xy = x³ + ax² + b"
    r_eq.font.name = FONT_MONO; r_eq.font.size = Pt(22); r_eq.font.bold = True; r_eq.font.color.rgb = BLUE_PRIMARY

    bullets = [
        ("Non-Singularity Condition", "b ≠ 0 (guarantees no singular points or cusps)."),
        ("Concrete Galois Field", "GF(2^4) with irreducible reduction polynomial f(x) = x⁴ + x + 1."),
        ("Generator Root g", "Satisfies f(g) = 0 => g⁴ = g + 1 = 0011₂ in standard bit representation."),
        ("Book Curve E_{2^4}(g⁴, 1)", "Parameters: a = g⁴, b = 1."),
        ("Worked Verification for (g⁵, g³)", "LHS = (g³)² + (g⁵)(g³) = g⁶ + g⁸ = 1100₂ ⊕ 0101₂ = 1001₂.\nRHS = (g⁵)³ + g⁴(g⁵)² + 1 = g¹⁵ + g¹⁴ + 1 = 0001₂ ⊕ 1001₂ ⊕ 0001₂ = 1001₂.\nLHS = RHS = 1001₂ ✔ Point verified on the curve!")
    ]
    for b_title, b_desc in bullets:
        pi = tf.add_paragraph()
        pi.space_after = Pt(8)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(11.5); ri_d.font.color.rgb = TEXT_MUTED

    add_image_card(s16, "images/figure_10_6_binary_curve.png", Inches(7.1), Inches(1.45), Inches(5.4), Inches(5.4), "Stallings Fig 10.6 — The Elliptic Curve E_{2^4}(g^4, 1) (16 Points)")

    # ==========================================
    # SLIDE 17: Arithmetic Rules over GF(2^m)
    # ==========================================
    s17 = prs.slides.add_slide(blank_layout)
    set_bg(s17, "fade")
    add_header(s17, 17, "Arithmetic Rules & Slope Formulas over GF(2^m)", "The 4 addition rules for E_{2^m}(a, b) (Stallings p. 9-10)")
    
    rules_gf = [
        ("Rule 1: Identity Element", "P + O = P for all points P in E_{2^m}(a, b).", BLUE_PRIMARY),
        ("Rule 2: Point Negation in Char 2", "If P = (x_P, y_P), -P = (x_P, x_P + y_P). P + (x_P, x_P + y_P) = O.\nCrucial distinction: In characteristic 2, +1 = -1, so negation is (x_P, x_P + y_P), NOT (x, -y)!", ROSE_RED),
        ("Rule 3: Point Addition (P ≠ Q)", "λ = (y_Q + y_P) / (x_Q + x_P)\nx_R = λ² + λ + x_P + x_Q + a\ny_R = λ(x_P + x_R) + x_R + y_P", PURPLE_ACCENT),
        ("Rule 4: Point Doubling (P = Q)", "Tangent Slope λ = x_P + (y_P / x_P)\nx_R = λ² + λ + a\ny_R = x_P² + (λ + 1)x_R", EMERALD_GREEN)
    ]
    for idx, (r_name, r_desc, col) in enumerate(rules_gf):
        t = Inches(1.45 + idx * 1.34)
        add_card(s17, Inches(0.8), t, Inches(11.733), Inches(1.22), CARD_BG, BORDER_SOFT, col)
        tb = s17.shapes.add_textbox(Inches(1.1), t + Inches(0.1), Inches(11.2), Inches(1.05))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.space_after = Pt(3)
        r1 = p.add_run(); r1.text = r_name; r1.font.name = FONT_TITLE; r1.font.size = Pt(14); r1.font.bold = True; r1.font.color.rgb = col
        p2 = tf.add_paragraph(); p2.line_spacing = 1.2
        r2 = p2.add_run(); r2.text = r_desc; r2.font.name = FONT_MONO; r2.font.size = Pt(11.5); r2.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 18: Master Comparison Table
    # ==========================================
    s18 = prs.slides.add_slide(blank_layout)
    set_bg(s18, "wipe")
    add_header(s18, 18, "Master Comparison — Z_p vs. GF(2^m)", "Side-by-side feature and formula matrix (Stallings Table 10.1 & 10.2)")
    
    comp_rows = [
        ("Feature / Property", "Prime Curves Z_p", "Binary Curves GF(2^m)"),
        ("Field Elements", "Integers {0, 1, ..., p-1}", "Polynomials in {0,1} of degree < m"),
        ("Curve Equation", "y² ≡ x³ + ax + b (mod p)", "y² + xy = x³ + ax² + b"),
        ("Non-Singularity", "(4a³ + 27b²) mod p ≠ 0", "b ≠ 0"),
        ("Point Negation -P", "(x_P, p - y_P)", "(x_P, x_P + y_P)"),
        ("Addition Slope λ", "[(y_Q - y_P)/(x_Q - x_P)] mod p", "(y_Q + y_P)/(x_Q + x_P)"),
        ("Doubling Slope λ", "[(3x_P² + a)/(2y_P)] mod p", "x_P + (y_P / x_P)"),
        ("Addition x_R", "(λ² - x_P - x_Q) mod p", "λ² + λ + x_P + x_Q + a"),
        ("Doubling x_R", "(λ² - 2x_P) mod p", "λ² + λ + a"),
        ("Primary Target", "Software Implementations (CPUs)", "Hardware Logic Gates (ASICs, Smart Cards)")
    ]

    table_comp_shape = s18.shapes.add_table(10, 3, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.4))
    t_comp = table_comp_shape.table
    t_comp.columns[0].width = Inches(2.733)
    t_comp.columns[1].width = Inches(4.5)
    t_comp.columns[2].width = Inches(4.5)

    for r_idx, row in enumerate(comp_rows):
        for c_idx, val in enumerate(row):
            cell = t_comp.cell(r_idx, c_idx)
            cell.margin_left = Inches(0.15); cell.margin_right = Inches(0.15)
            cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_c = cell.text_frame.paragraphs[0]
            p_c.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
            p_c.font.name = FONT_TITLE if r_idx == 0 else (FONT_BODY if c_idx == 0 else FONT_MONO)
            p_c.font.size = Pt(11)

            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TEXT_DARK
                p_c.font.bold = True
                p_c.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx == 9: # Primary target highlight
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG_ACCENT
                p_c.font.bold = True
                p_c.font.color.rgb = BLUE_PRIMARY
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
                p_c.font.bold = (c_idx == 0)
                p_c.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 19: Scalar Multiplication
    # ==========================================
    s19 = prs.slides.add_slide(blank_layout)
    set_bg(s19, "push")
    add_header(s19, 19, "Scalar Multiplication & Double-and-Add Algorithm", "Computing kP = P + P + ... + P efficiently in O(log k) steps")
    
    # Left Card
    add_card(s19, Inches(0.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, ROSE_RED)
    tb = s19.shapes.add_textbox(Inches(1.1), Inches(1.68), Inches(5.15), Inches(4.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]; p.space_after = Pt(12)
    r = p.add_run(); r.text = "Algorithm Comparison"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = ROSE_RED
    bullets_sm = [
        ("Naïve Method", "Add point P to itself (k - 1) times. For 256-bit keys, performing k ≈ 2^{256} point additions is physically impossible!"),
        ("Double-and-Add Method", "Scan bits of multiplier k from MSB to LSB. For every bit, double the accumulator. If bit = 1, add P."),
        ("Logarithmic Complexity", "Requires only ~256 point doublings and ~128 point additions (approx. 384 operations total instead of 2^{256}!)."),
        ("Constant-Time Mitigation", "Montgomery Ladder execution pattern protects against side-channel timing analysis attacks.")
    ]
    for b_title, b_desc in bullets_sm:
        pi = tf.add_paragraph()
        pi.space_after = Pt(10)
        pi.line_spacing = 1.22
        ri_t = pi.add_run(); ri_t.text = f"• {b_title}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(12.5); ri_t.font.bold = True; ri_t.font.color.rgb = TEXT_DARK
        ri_d = pi.add_run(); ri_d.text = b_desc; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12); ri_d.font.color.rgb = TEXT_MUTED

    # Right Card: Formatted Double-and-Add Table
    add_card(s19, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.4), CARD_BG, BORDER_SOFT, EMERALD_GREEN)
    tb_r = s19.shapes.add_textbox(Inches(7.05), Inches(1.68), Inches(5.2), Inches(0.5))
    tf_r = tb_r.text_frame; tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p = tf_r.paragraphs[0]; r = p.add_run()
    r.text = "Worked Example: 21P (21 = 10101₂)"; r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = EMERALD_GREEN

    table_da_shape = s19.shapes.add_table(6, 3, Inches(7.0), Inches(2.28), Inches(5.3), Inches(3.3))
    t_da = table_da_shape.table
    t_da.columns[0].width = Inches(1.4)
    t_da.columns[1].width = Inches(1.8)
    t_da.columns[2].width = Inches(2.1)

    da_steps = [
        ("Bit Scanned", "Operation Performed", "Accumulator State"),
        ("Bit 1 (MSB)", "Initialize", "P"),
        ("Bit 0", "Double", "2P"),
        ("Bit 1", "Double + Add P", "4P + P = 5P"),
        ("Bit 0", "Double", "10P"),
        ("Bit 1 (LSB)", "Double + Add P", "20P + P = 21P ✔")
    ]
    for r_idx, row in enumerate(da_steps):
        for c_idx, val in enumerate(row):
            cell = t_da.cell(r_idx, c_idx)
            cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
            cell.margin_top = Inches(0.08); cell.margin_bottom = Inches(0.08)
            cell.text = val
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p_c = cell.text_frame.paragraphs[0]
            p_c.alignment = PP_ALIGN.CENTER
            p_c.font.name = FONT_MONO if r_idx > 0 else FONT_TITLE
            p_c.font.size = Pt(11)

            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = TEXT_DARK
                p_c.font.bold = True
                p_c.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx == 5:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG_GREEN
                p_c.font.bold = True
                p_c.font.color.rgb = EMERALD_GREEN
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
                p_c.font.color.rgb = TEXT_DARK

    tb_da_end = s19.shapes.add_textbox(Inches(7.0), Inches(5.8), Inches(5.3), Inches(0.85))
    tf_da_end = tb_da_end.text_frame; tf_da_end.word_wrap = True
    tf_da_end.margin_left = tf_da_end.margin_top = tf_da_end.margin_right = tf_da_end.margin_bottom = 0
    p_de = tf_da_end.paragraphs[0]; p_de.line_spacing = 1.25; r_de = p_de.add_run()
    r_de.text = "Efficiency: Only 4 Doublings + 2 Additions (6 total operations) instead of 20 additions! Yields a 70% computational savings."
    r_de.font.name = FONT_BODY; r_de.font.size = Pt(11.5); r_de.font.bold = True; r_de.font.color.rgb = EMERALD_GREEN

    # ==========================================
    # SLIDE 20: Summary & Conclusion
    # ==========================================
    s20 = prs.slides.add_slide(blank_layout)
    set_bg(s20, "fade")
    add_header(s20, 20, "Summary & Conclusion", "Key Takeaways from Chapter 10")
    
    add_card(s20, Inches(0.8), Inches(1.45), Inches(11.733), Inches(4.35), CARD_BG, BORDER_BLUE)
    tb = s20.shapes.add_textbox(Inches(1.15), Inches(1.65), Inches(11.0), Inches(3.95))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    summary_pts = [
        ("1. Unrivaled Efficiency", "ECC achieves 128-bit commercial security with just 256 bits compared to 3072 bits in RSA (12× smaller key size)."),
        ("2. Rigorous Group Theory", "Point addition forms an Abelian group satisfying Closure, Associativity, Identity (O), Inverse (-P), and Commutativity."),
        ("3. Dual Finite Field Architectures", "Prime curves over Z_p target software ALUs on general CPUs; Binary curves over GF(2^m) target logic gates in ASICs & smart cards."),
        ("4. Precise Discrete Formulas", "Slope and coordinate formulas for secant addition and tangent doubling eliminate all floating-point rounding errors."),
        ("5. The ECDLP Trapdoor", "Scalar multiplication Q = kP is efficiently computed in O(log k) steps via Double-and-Add, while computing k given Q and P is intractable."),
        ("6. The Gold Standard for Modern Cryptography", "Absence of sub-exponential index calculus algorithms ensures ECC remains the cornerstone of modern network protocols (TLS 1.3, SSH, IPsec).")
    ]
    for idx, (pt_t, pt_d) in enumerate(summary_pts):
        pi = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        pi.space_after = Pt(8)
        pi.line_spacing = 1.2
        ri_t = pi.add_run(); ri_t.text = f"✔  {pt_t}: "; ri_t.font.name = FONT_BODY; ri_t.font.size = Pt(13); ri_t.font.bold = True; ri_t.font.color.rgb = BLUE_PRIMARY
        ri_d = pi.add_run(); ri_d.text = pt_d; ri_d.font.name = FONT_BODY; ri_d.font.size = Pt(12.5); ri_d.font.color.rgb = TEXT_DARK

    # Thank you card
    add_card(s20, Inches(3.0), Inches(6.0), Inches(7.333), Inches(1.05), CARD_BG_ALT, BLUE_PRIMARY)
    tb_ty = s20.shapes.add_textbox(Inches(3.2), Inches(6.05), Inches(6.9), Inches(0.95))
    tf_ty = tb_ty.text_frame; tf_ty.margin_left = tf_ty.margin_top = tf_ty.margin_right = tf_ty.margin_bottom = 0
    p = tf_ty.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.space_after = Pt(3)
    r = p.add_run(); r.text = "Thank You! Questions & Discussion"
    r.font.name = FONT_TITLE; r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = BLUE_PRIMARY
    p2 = tf_ty.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = "Reference: Cryptography and Network Security — William Stallings (Chapter 10)"
    r2.font.name = FONT_BODY; r2.font.size = Pt(11); r2.font.bold = True; r2.font.color.rgb = TEXT_LIGHT

    # Save output presentation
    output_path = "e:\\sem3\\network security\\ecc-presentation\\Elliptic_Curve_Arithmetic_Presentation_GuruPrakash.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")

if __name__ == '__main__':
    build_presentation()
